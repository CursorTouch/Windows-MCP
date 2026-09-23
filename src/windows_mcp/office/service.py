"""Office document and SQLite service.

Third-party document imports are lazy so the base MCP can start without optional packages.
Every SQL statement that accepts user values should be parameterized with the parameters argument.
"""

from __future__ import annotations

import csv
import json
import sqlite3
from pathlib import Path
from typing import Any

_MAX_TEXT_RETURN = 200_000


def _as_bool(value: Any, default: bool = False) -> bool:
    """Normalize booleans passed as strings."""
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


def _load_json_or_scalar(value: Any) -> Any:
    """Parse a JSON string when possible, otherwise return the original value."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text:
        return value
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return value


def _as_matrix(value: Any) -> list[list[Any]]:
    """Normalize a scalar, one-dimensional, or two-dimensional value into rows."""
    parsed = _load_json_or_scalar(value)
    if parsed is None:
        return []
    if isinstance(parsed, tuple):
        parsed = list(parsed)
    if not isinstance(parsed, list):
        return [[parsed]]
    if not parsed:
        return []
    first = parsed[0]
    if isinstance(first, (list, tuple)):
        return [list(row) if isinstance(row, (list, tuple)) else [row] for row in parsed]
    return [parsed]


def _as_string_list(value: Any, name: str) -> list[str]:
    """Normalize a string, JSON array, or list into a list of strings."""
    parsed = _load_json_or_scalar(value)
    if parsed is None or parsed == "":
        return []
    if isinstance(parsed, (str, int, float)):
        return [part.strip() for part in str(parsed).split(",") if part.strip()]
    if isinstance(parsed, tuple):
        parsed = list(parsed)
    if isinstance(parsed, list):
        return [str(item).strip() for item in parsed if str(item).strip()]
    raise ValueError(f"{name} must be a string or list")


def _require_existing_file(path: str, suffixes: tuple[str, ...] = ()) -> Path:
    """Resolve and verify a readable file."""
    resolved = Path(path).expanduser().resolve()
    if not resolved.exists():
        raise FileNotFoundError(f"File not found: {resolved}")
    if not resolved.is_file():
        raise ValueError(f"Path is not a file: {resolved}")
    if suffixes and resolved.suffix.lower() not in suffixes:
        expected = ", ".join(suffixes)
        raise ValueError(f"Unsupported file extension {resolved.suffix!r}; expected {expected}")
    return resolved


def _prepare_output(path: str, overwrite: bool, suffixes: tuple[str, ...] = ()) -> Path:
    """Resolve a new output path and verify its extension and overwrite policy."""
    resolved = Path(path).expanduser().resolve()
    if suffixes and resolved.suffix.lower() not in suffixes:
        expected = ", ".join(suffixes)
        raise ValueError(f"Unsupported file extension {resolved.suffix!r}; expected {expected}")
    if resolved.exists() and resolved.is_dir():
        raise ValueError(f"Destination is a directory: {resolved}")
    if resolved.exists() and not overwrite:
        raise FileExistsError(f"File already exists: {resolved}. Set overwrite=True to replace it.")
    resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


def _require_openpyxl():
    """Import openpyxl or raise an actionable dependency error."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required for Excel operations. Install openpyxl>=3.1.5."
        ) from exc
    return openpyxl


def _workbook(path: str, read_only: bool = False, data_only: bool = False):
    """Load an .xlsx/.xlsm workbook."""
    source = _require_existing_file(path, (".xlsx", ".xlsm"))
    openpyxl = _require_openpyxl()
    return openpyxl.load_workbook(
        source,
        read_only=read_only,
        data_only=data_only,
        keep_vba=source.suffix.lower() == ".xlsm",
    )


def _worksheet(workbook: Any, sheet: str | None) -> Any:
    """Select a worksheet by name or return the active sheet."""
    if not sheet:
        return workbook.active
    if sheet not in workbook.sheetnames:
        raise ValueError(f"Worksheet not found: {sheet}")
    return workbook[sheet]


def _save_workbook(workbook: Any, path: str | Path) -> None:
    """Save a workbook to an explicit path."""
    workbook.save(str(path))


def excel_create(path: str, sheet_name: str = "Sheet", overwrite: bool = False) -> str:
    """Create a new .xlsx workbook without requiring Microsoft Excel."""
    openpyxl = _require_openpyxl()
    destination = _prepare_output(path, _as_bool(overwrite), (".xlsx",))
    workbook = openpyxl.Workbook()
    workbook.active.title = sheet_name or "Sheet"
    _save_workbook(workbook, destination)
    return f"Excel workbook created: {destination}\nSheet: {workbook.active.title}"


def excel_list_sheets(path: str) -> str:
    """List worksheets and dimensions in an .xlsx/.xlsm workbook."""
    workbook = _workbook(path, read_only=True)
    try:
        rows = [
            {
                "name": name,
                "max_row": workbook[name].max_row,
                "max_column": workbook[name].max_column,
            }
            for name in workbook.sheetnames
        ]
        return json.dumps(rows, ensure_ascii=False, indent=2)
    finally:
        workbook.close()


def excel_read(
    path: str,
    sheet: str | None = None,
    cell_range: str | None = None,
    data_only: bool = True,
) -> str:
    """Read a cell or rectangular range and return a JSON two-dimensional array."""
    workbook = _workbook(path, data_only=_as_bool(data_only, True))
    try:
        worksheet = _worksheet(workbook, sheet)
        target = worksheet[cell_range] if cell_range else worksheet
        if not isinstance(target, tuple):
            return json.dumps([[target.value]], ensure_ascii=False, default=str)
        rows: list[list[Any]] = []
        for row in target:
            if isinstance(row, tuple):
                rows.append([cell.value for cell in row])
            else:
                rows.append([row.value])
        return json.dumps(rows, ensure_ascii=False, default=str)
    finally:
        workbook.close()


def excel_write(
    path: str,
    values: Any,
    sheet: str | None = None,
    cell_range: str | None = None,
    start_cell: str = "A1",
) -> str:
    """Write a scalar, row, or two-dimensional array to an .xlsx workbook."""
    openpyxl = _require_openpyxl()
    from openpyxl.utils.cell import range_boundaries

    source = _require_existing_file(path, (".xlsx", ".xlsm"))
    workbook = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        worksheet = _worksheet(workbook, sheet)
        rows = _as_matrix(values)
        if not rows:
            return "Error: values cannot be empty."
        if cell_range:
            min_col, min_row, _, _ = range_boundaries(cell_range)
        else:
            try:
                min_col, min_row = openpyxl.utils.cell.coordinate_to_tuple(start_cell)
            except Exception as exc:
                raise ValueError(f"Invalid start_cell: {start_cell}") from exc

        count = 0
        for row_offset, row in enumerate(rows):
            for col_offset, value in enumerate(row):
                worksheet.cell(
                    row=min_row + row_offset,
                    column=min_col + col_offset,
                    value=value,
                )
                count += 1
        _save_workbook(workbook, source)
        return (
            f"Excel cells written: {count}\nFile: {source}\nSheet: {worksheet.title}\n"
            f"Start: row {min_row}, column {min_col}"
        )
    finally:
        workbook.close()


def excel_append(path: str, values: Any, sheet: str | None = None) -> str:
    """Append one or more rows to an .xlsx worksheet."""
    openpyxl = _require_openpyxl()
    source = _require_existing_file(path, (".xlsx", ".xlsm"))
    workbook = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        worksheet = _worksheet(workbook, sheet)
        rows = _as_matrix(values)
        if not rows:
            return "Error: values cannot be empty."
        start_row = worksheet.max_row + 1
        for row in rows:
            worksheet.append(row)
        _save_workbook(workbook, source)
        return (
            f"Excel rows appended: {len(rows)} (starting row {start_row})\n"
            f"File: {source}\nSheet: {worksheet.title}"
        )
    finally:
        workbook.close()


def excel_add_sheet(path: str, name: str) -> str:
    """Add a worksheet to an existing workbook."""
    openpyxl = _require_openpyxl()
    source = _require_existing_file(path, (".xlsx", ".xlsm"))
    workbook = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        if name in workbook.sheetnames:
            return f"Error: Worksheet already exists: {name}"
        workbook.create_sheet(title=name)
        _save_workbook(workbook, source)
        return f"Excel worksheet added: {name}\nFile: {source}"
    finally:
        workbook.close()


def excel_delete_sheet(path: str, name: str) -> str:
    """Delete a worksheet, preserving at least one worksheet in the workbook."""
    openpyxl = _require_openpyxl()
    source = _require_existing_file(path, (".xlsx", ".xlsm"))
    workbook = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        if name not in workbook.sheetnames:
            return f"Error: Worksheet not found: {name}"
        if len(workbook.sheetnames) == 1:
            return "Error: Cannot delete the only worksheet."
        workbook.remove(workbook[name])
        _save_workbook(workbook, source)
        return f"Excel worksheet deleted: {name}\nFile: {source}"
    finally:
        workbook.close()


def excel_set_formula(path: str, cell: str, formula: str, sheet: str | None = None) -> str:
    """Write a formula to a cell; formulas are stored as text such as =SUM(A1:A3)."""
    openpyxl = _require_openpyxl()
    source = _require_existing_file(path, (".xlsx", ".xlsm"))
    workbook = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        worksheet = _worksheet(workbook, sheet)
        if not formula.startswith("="):
            formula = f"={formula}"
        worksheet[cell] = formula
        _save_workbook(workbook, source)
        return f"Excel formula written: {cell} = {formula}\nSheet: {worksheet.title}"
    finally:
        workbook.close()


def excel_iterate_rows(
    path: str,
    sheet: str | None = None,
    min_row: int = 1,
    max_row: int | None = None,
    values_only: bool = True,
    limit: int = 1000,
) -> str:
    """Iterate worksheet rows and return JSON rows, optionally with cell objects."""
    workbook = _workbook(path)
    try:
        worksheet = _worksheet(workbook, sheet)
        rows: list[list[Any]] = []
        for row in worksheet.iter_rows(
            min_row=max(1, int(min_row)),
            max_row=int(max_row) if max_row not in (None, "") else None,
        ):
            if _as_bool(values_only, True):
                item = [cell.value for cell in row]
            else:
                item = [{"coordinate": cell.coordinate, "value": cell.value} for cell in row]
            rows.append(item)
            if len(rows) >= max(1, int(limit)):
                break
        return json.dumps(rows, ensure_ascii=False, default=str)
    finally:
        workbook.close()


def excel_save(path: str, destination: str | None = None, overwrite: bool = False) -> str:
    """Save an .xlsx workbook in place or to another path."""
    openpyxl = _require_openpyxl()
    source = _require_existing_file(path, (".xlsx", ".xlsm"))
    workbook = openpyxl.load_workbook(source, keep_vba=source.suffix.lower() == ".xlsm")
    try:
        if destination:
            target = _prepare_output(destination, _as_bool(overwrite), (".xlsx", ".xlsm"))
        else:
            target = source
        _save_workbook(workbook, target)
        return f"Excel workbook saved: {target}"
    finally:
        workbook.close()


def _document():
    """Import python-docx or raise an actionable dependency error."""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "python-docx is required for Word operations. Install python-docx>=1.2.0."
        ) from exc
    return Document


def word_create(path: str, overwrite: bool = False) -> str:
    """Create a .docx document without requiring Microsoft Word."""
    destination = _prepare_output(path, _as_bool(overwrite), (".docx",))
    document = _document()()
    document.save(str(destination))
    return f"Word document created: {destination}"


def word_read(path: str) -> str:
    """Read all paragraphs and table cell values from a .docx document."""
    source = _require_existing_file(path, (".docx",))
    document = _document()(str(source))
    paragraphs = [
        {
            "index": index,
            "style": paragraph.style.name if paragraph.style else "",
            "text": paragraph.text,
        }
        for index, paragraph in enumerate(document.paragraphs)
    ]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
    ]
    return json.dumps(
        {"path": str(source), "paragraphs": paragraphs, "tables": tables},
        ensure_ascii=False,
        indent=2,
    )


def word_write(
    path: str,
    text: str,
    append: bool = True,
    style: str | None = None,
    create_if_missing: bool = True,
) -> str:
    """Write a paragraph to a .docx document."""
    destination = Path(path).expanduser().resolve()
    if destination.suffix.lower() != ".docx":
        raise ValueError("Word documents require the .docx extension")
    Document = _document()
    if not destination.exists():
        if not _as_bool(create_if_missing, True):
            raise FileNotFoundError(f"File not found: {destination}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
    else:
        document = Document(str(destination))
    paragraph = document.add_paragraph(text)
    if style:
        paragraph.style = style
    document.save(str(destination))
    return f"Word paragraph added: {len(document.paragraphs)} total\nFile: {destination}"


def word_read_tables(path: str, table_index: int | None = None) -> str:
    """Read all Word tables or one zero-based table."""
    source = _require_existing_file(path, (".docx",))
    document = _document()(str(source))
    if table_index is not None:
        index = int(table_index)
        if index < 0 or index >= len(document.tables):
            return f"Error: table_index out of range: {index}"
        tables = [[cell.text for cell in row.cells] for row in document.tables[index].rows]
    else:
        tables = [
            [[cell.text for cell in row.cells] for row in table.rows] for table in document.tables
        ]
    return json.dumps(tables, ensure_ascii=False, indent=2)


def word_add_heading(path: str, text: str, level: int = 1) -> str:
    """Append a heading to a .docx document."""
    source = _require_existing_file(path, (".docx",))
    document = _document()(str(source))
    document.add_heading(text, level=max(1, min(9, int(level))))
    document.save(str(source))
    return f"Word heading added: {text!r}\nFile: {source}"


def word_add_picture(path: str, image_path: str, width_inches: float | None = None) -> str:
    """Append a picture to a .docx document."""
    source = _require_existing_file(path, (".docx",))
    image = _require_existing_file(
        image_path, (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff")
    )
    document = _document()(str(source))
    kwargs = {}
    if width_inches not in (None, ""):
        from docx.shared import Inches

        kwargs["width"] = Inches(float(width_inches))
    document.add_picture(str(image), **kwargs)
    document.save(str(source))
    return f"Word picture added: {image}\nFile: {source}"


def _pypdf():
    """Import pypdf or raise an actionable dependency error."""
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError as exc:
        raise RuntimeError("pypdf is required for PDF operations. Install pypdf>=6.19.0.") from exc
    return PdfReader, PdfWriter


def _parse_pages(value: Any, page_count: int) -> list[int]:
    """Parse 1-based page selectors such as 1,3-5 into zero-based indexes."""
    if value in (None, "", []):
        return list(range(page_count))
    parsed = _load_json_or_scalar(value)
    if isinstance(parsed, int):
        raw_items: list[Any] = [parsed]
    elif isinstance(parsed, str):
        raw_items = [part.strip() for part in parsed.split(",") if part.strip()]
    elif isinstance(parsed, (list, tuple)):
        raw_items = list(parsed)
    else:
        raise ValueError("pages must be an integer, comma-separated string, or list")
    indexes: list[int] = []
    for item in raw_items:
        text = str(item).strip()
        if "-" in text:
            start_text, end_text = text.split("-", 1)
            start, end = int(start_text), int(end_text)
            if end < start:
                start, end = end, start
            candidates = range(start, end + 1)
        else:
            candidates = [int(text)]
        for page in candidates:
            index = page - 1
            if index < 0 or index >= page_count:
                raise ValueError(f"Page {page} is out of range; document has {page_count} pages")
            if index not in indexes:
                indexes.append(index)
    return indexes


def _reader_for(path: str, password: str | None = None):
    """Open a PDF reader and decrypt it when needed."""
    PdfReader, _ = _pypdf()
    reader = PdfReader(str(_require_existing_file(path, (".pdf",))))
    if reader.is_encrypted:
        result = reader.decrypt(password or "")
        if result == 0:
            raise ValueError(f"Unable to decrypt PDF with the supplied password: {path}")
    return reader


def pdf_info(path: str, password: str | None = None) -> str:
    """Return PDF page count and metadata."""
    reader = _reader_for(path, password)
    metadata = {str(key): str(value) for key, value in dict(reader.metadata or {}).items()}
    return json.dumps(
        {
            "path": str(Path(path).expanduser().resolve()),
            "pages": len(reader.pages),
            "metadata": metadata,
        },
        ensure_ascii=False,
        indent=2,
    )


def pdf_extract_text(
    path: str,
    pages: Any = None,
    password: str | None = None,
    max_chars: int = _MAX_TEXT_RETURN,
) -> str:
    """Extract text page by page from a PDF."""
    reader = _reader_for(path, password)
    selected = _parse_pages(pages, len(reader.pages))
    limit = max(0, int(max_chars))
    result = []
    for index in selected:
        text = reader.pages[index].extract_text() or ""
        result.append({"page": index + 1, "text": text[:limit], "truncated": len(text) > limit})
    return json.dumps(result, ensure_ascii=False, indent=2)


def pdf_merge(
    sources: str | list[str],
    destination: str,
    overwrite: bool = False,
    password: str | None = None,
) -> str:
    """Merge multiple PDFs into a new PDF."""
    source_values = _as_string_list(sources, "sources")
    if len(source_values) < 2:
        return "Error: At least two PDF sources are required for merge."
    target = _prepare_output(destination, _as_bool(overwrite), (".pdf",))
    source_paths = [_require_existing_file(item, (".pdf",)) for item in source_values]
    if target in source_paths:
        return "Error: destination must differ from every source PDF."
    _, PdfWriter = _pypdf()
    writer = PdfWriter()
    page_count = 0
    try:
        for source in source_paths:
            reader = _reader_for(str(source), password)
            for page in reader.pages:
                writer.add_page(page)
                page_count += 1
        with target.open("wb") as handle:
            writer.write(handle)
    finally:
        writer.close()
    return f"PDF merged: {len(source_paths)} files, {page_count} pages\nFile: {target}"


def pdf_extract_pages(
    path: str,
    destination: str,
    pages: Any,
    overwrite: bool = False,
    password: str | None = None,
) -> str:
    """Extract selected pages from a PDF into a new PDF."""
    reader = _reader_for(path, password)
    selected = _parse_pages(pages, len(reader.pages))
    if not selected:
        return "Error: No pages selected."
    target = _prepare_output(destination, _as_bool(overwrite), (".pdf",))
    _, PdfWriter = _pypdf()
    writer = PdfWriter()
    try:
        for index in selected:
            writer.add_page(reader.pages[index])
        with target.open("wb") as handle:
            writer.write(handle)
    finally:
        writer.close()
    return f"PDF pages extracted: {len(selected)}\nFile: {target}"


def pdf_split(
    path: str,
    destination_dir: str,
    overwrite: bool = False,
    password: str | None = None,
) -> str:
    """Split every PDF page into a separate file inside a destination directory."""
    reader = _reader_for(path, password)
    source = Path(path).expanduser().resolve()
    directory = Path(destination_dir).expanduser().resolve()
    directory.mkdir(parents=True, exist_ok=True)
    _, PdfWriter = _pypdf()
    outputs: list[str] = []
    for index, page in enumerate(reader.pages):
        target = directory / f"{source.stem}_page_{index + 1:03d}.pdf"
        if target.exists() and not _as_bool(overwrite):
            raise FileExistsError(f"File already exists: {target}. Set overwrite=True.")
        writer = PdfWriter()
        try:
            writer.add_page(page)
            with target.open("wb") as handle:
                writer.write(handle)
        finally:
            writer.close()
        outputs.append(str(target))
    return f"PDF split into {len(outputs)} files:\n" + "\n".join(outputs)


def _detect_encoding(path: Path) -> str:
    """Detect common text encodings without a third-party dependency."""
    raw = path.read_bytes()[:131072]
    if raw.startswith(b"\xef\xbb\xbf"):
        return "utf-8-sig"
    for encoding in ("utf-8", "gb18030", "cp1252"):
        try:
            raw.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "utf-8"


def _detect_delimiter(sample: str, requested: str | None = None) -> str:
    """Detect a CSV delimiter with csv.Sniffer and a safe fallback."""
    if requested:
        if len(requested) != 1:
            raise ValueError("delimiter must be one character")
        return requested
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def csv_read(
    path: str,
    delimiter: str | None = None,
    encoding: str | None = None,
    has_header: bool = True,
    max_rows: int = 1000,
) -> str:
    """Read CSV/TSV data with automatic encoding and delimiter detection."""
    source = _require_existing_file(path, (".csv", ".tsv", ".txt"))
    actual_encoding = encoding or _detect_encoding(source)
    with source.open("r", encoding=actual_encoding, errors="replace", newline="") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        actual_delimiter = _detect_delimiter(sample, delimiter)
        reader = csv.reader(handle, delimiter=actual_delimiter)
        raw_rows = list(reader)
    limit = max(1, int(max_rows))
    if _as_bool(has_header, True) and raw_rows:
        header = raw_rows[0]
        rows = raw_rows[1 : limit + 1]
    else:
        header = []
        rows = raw_rows[:limit]
    data_row_count = len(raw_rows) - (1 if header else 0)
    return json.dumps(
        {
            "path": str(source),
            "encoding": actual_encoding,
            "delimiter": actual_delimiter,
            "header": header,
            "rows": rows,
            "row_count": data_row_count,
            "truncated": data_row_count > len(rows),
        },
        ensure_ascii=False,
        indent=2,
    )


def csv_write(
    path: str,
    rows: Any,
    delimiter: str = ",",
    encoding: str = "utf-8-sig",
    overwrite: bool = False,
) -> str:
    """Write a two-dimensional array to CSV with UTF-8 BOM by default."""
    matrix = _as_matrix(rows)
    if not matrix:
        return "Error: rows cannot be empty."
    if len(delimiter) != 1:
        return "Error: delimiter must be one character."
    target = _prepare_output(path, _as_bool(overwrite), (".csv", ".tsv", ".txt"))
    with target.open("w", encoding=encoding, newline="") as handle:
        writer = csv.writer(handle, delimiter=delimiter)
        writer.writerows(matrix)
    return (
        f"CSV written: {target}\nRows: {len(matrix)}\n"
        f"Delimiter: {delimiter!r}\nEncoding: {encoding}"
    )


def _sql_parameters(value: Any) -> Any:
    """Normalize JSON string parameters into a sqlite3-compatible sequence/mapping."""
    parsed = _load_json_or_scalar(value)
    if parsed is None or parsed == "":
        return []
    if isinstance(parsed, (list, tuple, dict)):
        return parsed
    return [parsed]


def _require_sql(path: str, create: bool = False) -> Path:
    """Resolve a SQLite path and enforce existence unless creating."""
    resolved = Path(path).expanduser().resolve()
    if not create and not resolved.exists():
        raise FileNotFoundError(f"Database not found: {resolved}")
    if resolved.exists() and resolved.is_dir():
        raise ValueError(f"Database path is a directory: {resolved}")
    resolved.parent.mkdir(parents=True, exist_ok=True)
    return resolved


def sqlite_create(path: str, overwrite: bool = False) -> str:
    """Create an empty SQLite database."""
    target = _prepare_output(path, _as_bool(overwrite), (".db", ".sqlite", ".sqlite3"))
    connection = sqlite3.connect(target)
    try:
        connection.execute("PRAGMA journal_mode=WAL")
        connection.commit()
    finally:
        connection.close()
    return f"SQLite database created: {target}"


def sqlite_execute(
    path: str,
    sql: str,
    parameters: Any = None,
    create: bool = False,
) -> str:
    """Execute one parameterized DML/DDL statement.

    Always pass user values through parameters; never interpolate untrusted text into SQL.
    """
    target = _require_sql(path, _as_bool(create))
    connection = sqlite3.connect(target)
    try:
        cursor = connection.execute(sql, _sql_parameters(parameters))
        connection.commit()
        return json.dumps(
            {"path": str(target), "rowcount": cursor.rowcount, "lastrowid": cursor.lastrowid},
            ensure_ascii=False,
            indent=2,
        )
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def sqlite_query(
    path: str,
    sql: str,
    parameters: Any = None,
    limit: int | None = None,
) -> str:
    """Run a parameterized SELECT/PRAGMA query and return named JSON rows."""
    target = _require_sql(path)
    connection = sqlite3.connect(target)
    try:
        cursor = connection.execute(sql, _sql_parameters(parameters))
        if not cursor.description:
            return "Error: Query did not return columns."
        columns = [item[0] for item in cursor.description]
        if limit not in (None, ""):
            rows = cursor.fetchmany(max(0, int(limit)))
        else:
            rows = cursor.fetchall()
        return json.dumps(
            {"columns": columns, "rows": [dict(zip(columns, row)) for row in rows]},
            ensure_ascii=False,
            default=str,
            indent=2,
        )
    finally:
        connection.close()


def _quote_identifier(identifier: str) -> str:
    """Quote a SQLite identifier while preserving it as a single name."""
    if not identifier or "\x00" in identifier:
        raise ValueError("Invalid empty or NUL-containing SQL identifier")
    return '"' + identifier.replace('"', '""') + '"'


def sqlite_batch_insert(
    path: str,
    table: str,
    columns: str | list[str],
    rows: Any,
    create: bool = False,
) -> str:
    """Insert rows efficiently with executemany and bound parameters."""
    column_list = _as_string_list(columns, "columns")
    if not column_list:
        return "Error: columns cannot be empty."
    matrix = _as_matrix(rows)
    if not matrix:
        return "Error: rows cannot be empty."
    width = len(column_list)
    for index, row in enumerate(matrix, start=1):
        if len(row) != width:
            return f"Error: row {index} has {len(row)} values; expected {width}."
    target = _require_sql(path, _as_bool(create))
    quoted_table = _quote_identifier(table)
    quoted_columns = ", ".join(_quote_identifier(column) for column in column_list)
    placeholders = ", ".join("?" for _ in column_list)
    statement = f"INSERT INTO {quoted_table} ({quoted_columns}) VALUES ({placeholders})"
    connection = sqlite3.connect(target)
    try:
        cursor = connection.executemany(statement, matrix)
        connection.commit()
        return json.dumps(
            {"path": str(target), "inserted": cursor.rowcount, "table": table},
            ensure_ascii=False,
            indent=2,
        )
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def sqlite_transaction(path: str, statements: Any, create: bool = False) -> str:
    """Execute a JSON list of statements atomically and rollback on any error.

    Each statement may be a SQL string, [sql, parameters], or an object with sql and
    parameters fields. Use placeholders and parameter values for untrusted data.
    """
    parsed = _load_json_or_scalar(statements)
    if not isinstance(parsed, list) or not parsed:
        return "Error: statements must be a non-empty JSON array."
    target = _require_sql(path, _as_bool(create))
    connection = sqlite3.connect(target)
    try:
        connection.execute("BEGIN IMMEDIATE")
        count = 0
        for item in parsed:
            if isinstance(item, str):
                sql, params = item, []
            elif isinstance(item, dict):
                sql, params = item.get("sql"), item.get("parameters", [])
            elif isinstance(item, (list, tuple)) and item:
                sql = item[0]
                params = item[1] if len(item) > 1 else []
            else:
                raise ValueError(f"Invalid transaction statement: {item!r}")
            if not isinstance(sql, str) or not sql.strip():
                raise ValueError("Every transaction statement requires a non-empty SQL string.")
            connection.execute(sql, _sql_parameters(params))
            count += 1
        connection.commit()
        return json.dumps({"path": str(target), "statements_committed": count}, indent=2)
    except Exception:
        connection.rollback()
        raise
    finally:
        connection.close()


def sqlite_tables(path: str) -> str:
    """List user tables and views in a SQLite database."""
    return sqlite_query(
        path,
        "SELECT type, name, tbl_name, sql FROM sqlite_master "
        "WHERE type IN ('table', 'view') ORDER BY type, name",
    )
